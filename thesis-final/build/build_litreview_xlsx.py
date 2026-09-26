# Literature review as ONE simple Excel sheet: one row per paper.
# Our numbers come from ../../results/summary_all.csv; everything else is from the papers.
import csv, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.join(HERE, '..', '..')
OUT = os.path.join(HERE, '..', 'Literature_Review.xlsx')
RES = {r['run']: r for r in csv.DictReader(open(os.path.join(ROOT, 'results', 'summary_all.csv')))}
F = lambda r: RES[r]['FND']

NAVY, TINT, HL, GROUP = '14213D', 'F3F6FB', 'DCE9FA', 'E8EEF7'
font = lambda **k: Font(name='Arial', **k)
thin = Side(style='thin', color='D0CFC6'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

# Paper, Year, Idea, Reported result, Weakness, What we took, Our result, Reference
ROWS = [
 ('Classic protocols', None),
 ('LEACH', 2000, 'Random CH rotation: threshold T(n) + epoch set G (a node that served waits until all have served). Members join the nearest CH; the CH fuses the data.',
  'FND 932, LND 1312 rounds (0.5 J).', 'Ignores energy: a weak or badly placed node can become CH.', 'The rotation T(r), the G-set and the nearest-CH join.',
  f"FND {F('leach_EDITED')}", '[1] Heinzelman et al., Proc. HICSS, 2000'),
 ('LEACH-C', 2002, 'The BS chooses the CHs centrally among nodes with at least the average energy; defines the two-slope radio model.',
  'Better CH placement than LEACH.', 'Needs the position and energy of every node at the BS every round.', 'The radio model; the "≥ average energy" idea (I5).',
  'radio model only', '[2] Heinzelman et al., IEEE Trans. Wireless Commun., 2002'),
 ('HEED', 2004, 'Residual energy decides who leads: CH_prob = C_prob·E_res/E_max, doubled every iteration; nodes negotiate until final CHs are known.',
  'Longer lifetime than LEACH (graphs only).', 'The negotiation messages are paid every round.', 'Energy × neighbours — but computed in ONE step, no negotiation.',
  f"FND {F('heed_EDITED')}", '[3] Younis & Fahmy, IEEE Trans. Mobile Comput., 2004'),
 ('PEGASIS', 2002, 'One greedy chain; each node fuses and passes the data to its neighbour; one leader (i mod N) sends to the BS.',
  'FND 1578, HND 2082, LND 2192 rounds.', 'Long delay; long chain links; the leader is a single point of failure.', 'Relaying between CHs — only when it saves energy (v8-Chain).',
  f"FND {F('pegasis_EDITED')}", '[4] Lindsey & Raghavendra, IEEE Aerospace Conf., 2002'),
 ('LEACH + HEED hybrids', None),
 ('SH-LEACH', 2015, 'LEACH draw with an energy-aware probability that also grows with the round number r.',
  'Later first death than LEACH (graphs only).', 'r is never reset → up to 100 CHs per round.', 'Lesson: keep the rotation. We fixed it (r restarts every 10 rounds).',
  f"FND {F('shleach_EDITED')} → fixed {F('shleach_IMPROVED')}", '[5] Shrestha et al., IJIBC, 2015'),
 ('H-LEACH', 2016, 'LEACH threshold with a personal energy-based probability; only nodes with E > average may become CH.',
  'LND ≈ 4312 vs 1778 for LEACH.', 'Equal energy at the start → nobody is above average → no CH (deadlock); no G-set.', 'The energy gate, written as E ≥ average (I5). We fixed it (≥ + G-set).',
  f"FND {F('hleach_EDITED')} → fixed {F('hleach_IMPROVED')}", '[6] Razaque et al., IEEE LISAT, 2016'),
 ('EECH-HEED', 2025, 'Two zones: HEED near the BS, energy × degree far away; relays between CHs; adaptive sensing.',
  'FND 1250, LND 2200, PDR 95 %.', 'Zone 2: P ≈ 1 → the rotation collapses (same nodes are CH every other round).', 'Energy × degree score; relaying between CHs. We fixed the rotation (10 %).',
  f"FND {F('eechheed_EDITED')} → fixed {F('eechheed_IMPROVED')}", '[7] Kaur et al., Scientific Reports, 2025'),
 ('Recent work (studied, not re-implemented)', None),
 ('RL-ILEACH', 2026, 'Q-learning chooses the CHs from residual energy, node density and distance.',
  '+784 % lifetime vs LEACH; FND 524 (200 nodes).', 'First node dies earlier than ILEACH (524 vs 936); learning cost not measured.', 'Shows that a closed-form score can reach a later first death.',
  '—', '[8] El-Sayed et al., Scientific Reports, 2026'),
 ('DL-HEED', 2025, 'A graph neural network inside HEED chooses the CHs (energy, degree, position, signal).',
  'Up to +60 % lifetime vs HEED.', 'Needs training data; heavy for sensor nodes.', 'Confirms that energy + degree are the key features.',
  '—', '[9] Juwaied & Jackowska-Strumillo, Applied Sciences, 2025'),
 ('TLC', 2024, 'Two-level LEACH-based CH selection with a distance threshold.',
  'Longer lifetime than LEACH variants.', 'Depends on the deployment geometry; no protection if a CH fails.', 'Motivates protecting the CH, not only placing it.',
  '—', '[10] Subedi et al., Telecom, 2024'),
 ('Ideas used (fault tolerance, rotation, relaying)', None),
 ('FTEC', 2019, 'Fault-tolerant clustering: backup mechanisms keep a cluster working when its CH fails.',
  'Clusters survive CH failures.', 'No energy-aware relaying between CHs.', 'The backup CH and cluster repair (v5).',
  '—', '[11] Jassbi & Moridi, Wireless Pers. Commun., 2019'),
 ('CH-rotation review', 2022, 'Survey of time-driven vs energy-driven CH rotation and their overhead.',
  '—', 'Review only.', 'Keep clusters for several rounds + hand over on low energy (I2).',
  '—', '[12] Pachlor et al., arXiv, 2022'),
 ('EEUC', 2005, 'Unequal clusters: smaller clusters near the BS so relaying CHs keep energy.',
  'Balances the relay load (no hot spot).', 'Needs tuning of the competition radius.', 'The "hot spot" lesson: relay only when it is affordable.',
  '—', '[13] Li et al., IEEE MASS, 2005'),
 ('Chauhan & Soni', 2021, 'Unequal clustering with multi-hop through low-degree relay nodes.',
  'Longer lifetime with relay nodes.', 'Extra route computation.', 'Relay only through CHs that can afford it.',
  '—', '[14] Chauhan & Soni, J. Ambient Intell. Humaniz. Comput., 2021'),
]
OURS = ('Our protocol: v8 / v8-Chain', 2026,
        'Single-pass score = LEACH turn × energy × neighbours; energy gate; clusters kept 5 rounds; handover, backup and re-join; direct-to-BS; relay between CHs only if cheaper (v8-Chain).',
        '—', 'Tested with one network size (100 nodes, 100 × 100 m); analytical radio model.', '—',
        f"FND {F('v8_center')} (far BS: v8-Chain {F('v8_chain_farBS')})", 'This thesis')

HEAD = ['Paper', 'Year', 'Idea (how it works)', 'Reported result (paper)', 'Weakness', 'What we took for our protocol', 'Our result (same environment, BS at the centre)', 'Reference']
WID = [18, 7, 46, 26, 36, 38, 22, 34]

wb = Workbook(); ws = wb.active; ws.title = 'Literature Review'; ws.sheet_view.showGridLines = False
ws['A1'] = 'Literature Review — clustering protocols for energy-efficient WSNs'; ws['A1'].font = font(bold=True, size=15, color=NAVY)
ws['A2'] = 'One row per paper. "Our result" = first node death (FND, rounds) of our ns-3.41 code in the same environment for all (100 nodes, 100 × 100 m, 0.5 J, BS at the centre).'
ws['A2'].font = font(italic=True, size=10, color='5F5E5A')
r = 4
for j, (h, w) in enumerate(zip(HEAD, WID), 1):
    c = ws.cell(r, j, h); c.font = font(bold=True, color='FFFFFF', size=11); c.fill = PatternFill('solid', fgColor=NAVY)
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = BOX
    ws.column_dimensions[get_column_letter(j)].width = w
ws.row_dimensions[r].height = 46
ws.freeze_panes = 'B5'

def row(vals, fill, bold=False):
    global r
    r += 1
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v); c.border = BOX; c.fill = PatternFill('solid', fgColor=fill)
        c.font = font(size=10, bold=bold or j == 1, color=NAVY if j == 1 else '1F1F1D')
        c.alignment = Alignment(wrap_text=True, vertical='top', horizontal='center' if j == 2 else 'left')
    longest = max(len(str(v)) / max(WID[j] - 2, 1) for j, v in enumerate(vals))
    ws.row_dimensions[r].height = max(30, 14 * (int(longest) + 1) + 4)

k = 0
for item in ROWS:
    if item[1] is None:
        r += 1
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(HEAD))
        c = ws.cell(r, 1, item[0]); c.font = font(bold=True, size=11, color=NAVY); c.fill = PatternFill('solid', fgColor=GROUP)
        c.alignment = Alignment(vertical='center'); ws.row_dimensions[r].height = 22
        continue
    row(item, TINT if k % 2 == 0 else 'FFFFFF'); k += 1
r += 1
row(OURS, HL, bold=True)

r += 2
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(HEAD))
c = ws.cell(r, 1, 'Research gap: no reviewed protocol (1) elects CHs in one step from energy and neighbours, (2) keeps clusters for several rounds, '
                  '(3) protects the CH before and after it fails, and (4) relays between CHs only when that saves energy — and they were never compared in one environment. '
                  'v8 / v8-Chain fills this gap.')
c.font = font(bold=True, size=11, color=NAVY); c.fill = PatternFill('solid', fgColor='FFF7E6')
c.alignment = Alignment(wrap_text=True, vertical='center'); c.border = BOX
ws.row_dimensions[r].height = 48
ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
wb.save(OUT); print('saved', OUT)
