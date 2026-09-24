import json, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
D=json.load(open('data_ar.json'))
RED='8A2522'; CHAR='323232'; CREAM='EDEDE5'; PINK='F6E3E1'
F=lambda **k: Font(name='Arial',**k)
thin=Side(style='thin',color='D0CFC6'); B=Border(bottom=thin)
HF=PatternFill('solid',fgColor=RED); BAND=PatternFill('solid',fgColor=CREAM); HL=PatternFill('solid',fgColor=PINK); NOF=PatternFill()
PCT='0.00%'; INT='#,##0'; DPCT='+0.0%;-0.0%;0.0%'; DPP='+0.00" pp";-0.00" pp";0.00" pp"'
R=lambda **k: Alignment(horizontal=k.get('h','right'),vertical='center',wrap_text=k.get('w',False),readingOrder=2)
S_README='اقرأني'; S_DASH='لوحة الملخص'; S_UNI='المقارنة الموحدة'; S_REP='إعادة البناء'; S_EVO='تطور البروتوكول'; S_CF='نص مقابل بعيد'; S_V8='تحليل v8'; S_GL='شرح المصطلحات'
q=lambda s: "'"+s+"'"
wb=openpyxl.Workbook()
def sheet(ws):
    ws.sheet_view.rightToLeft=True; ws.sheet_view.showGridLines=False
def title(ws,t,sub):
    sheet(ws); ws['A1']=t; ws['A1'].font=F(bold=True,size=16,color=RED); ws['A1'].alignment=R()
    ws['A2']=sub; ws['A2'].font=F(italic=True,size=10,color='555555'); ws['A2'].alignment=R()
def note(ws,r,txt,c=1):
    ws.cell(r,c,txt).font=F(italic=True,size=10,color='555555'); ws.cell(r,c).alignment=R()
def header(ws,r,cols,calc=()):
    for j,h in enumerate(cols,1):
        c=ws.cell(r,j,h); c.font=F(bold=True,color='FFFFFF',size=10); c.fill=HF if j not in calc else PatternFill('solid',fgColor=CHAR)
        c.alignment=R(h='center',w=True)
    ws.row_dimensions[r].height=34
def row(ws,r,vals,fmts,hl=False,band=False):
    for j,v in enumerate(vals,1):
        c=ws.cell(r,j,v); c.font=F(size=10,bold=hl); c.border=B
        c.fill=HL if hl else (BAND if band else NOF)
        if j in fmts: c.number_format=fmts[j]
        c.alignment=R(h='right' if isinstance(v,str) and not v.startswith('=') else 'center',w=isinstance(v,str) and len(v)>26)
def widths(ws,m):
    for col,w in m.items(): ws.column_dimensions[col].width=w

# ---------- README
ws=wb.active; ws.title=S_README
title(ws,'ملف مقارنة بروتوكولات الـ Clustering في شبكات الحساسات','رسالة: Energy-Efficient Secure Clustering in WSNs Using Hybrid Cryptography — م. رحمة خالد عشبة')
lines=[('الملف ده فيه إيه؟',None),
 (S_DASH,'أهم النتايج في صفحة واحدة: أحسن نسخة من كل عيلة بروتوكولات جنب بعض، ومعاها رسمين بيانيين.'),
 (S_UNI,'كل البروتوكولات وكل نسخ البروتوكول المقترح في نفس البيئة بالظبط (100 نود، 100×100 م، الـ BS في النص). الترتيب ونسب التغيير محسوبين بمعادلات.'),
 (S_REP,'أرقام الأبحاث الأصلية، مقابل إعادة بناءنا بإعدادات البحث نفسه، مقابل البيئة الموحّدة. والفرق بينهم محسوب بمعادلات.'),
 (S_EVO,'النسخ من v1 لحد v8-Chain، وكل نسخة بتغيّر حاجة واحدة بس. التغيير عن النسخة اللي قبلها وعن v1 محسوب بمعادلات.'),
 (S_CF,'نفس البروتوكول مرة والـ BS في نص الأرض (50,50)، ومرة والـ BS بعيد (50,−100).'),
 (S_V8,'تحليل البروتوكول النهائي: نشيل كل آلية لوحدها ونشوف أثرها، ونقارن طرق التوجيه، ونجرّب 8 أشكال شبكات مختلفة.'),
 (S_GL,'شرح كل مصطلح وكل عمود بالعربي البسيط.'),
 ('',None),('الألوان',None),
 ('العنوان الأحمر','رقم طالع من المحاكاة (اتكتب بإيدنا).'),
 ('العنوان الرمادي الغامق','عمود محسوب بمعادلة — ماتكتبيش فوقه.'),
 ('الصف الوردي','البروتوكول النهائي المقترح (v8 / v8-Chain).'),
 ('',None),('إعدادات المحاكاة',None)]+[(a,b) for a,b in D['env']]+[('',None),('المراجع',None)]+[(f'[{i}]',r) for i,r in enumerate(D['refs'],1)]
r=4
for a,b in lines:
    c=ws.cell(r,1,a); c.font=F(bold=True,size=11 if b is None else 10,color=RED if b is None else CHAR); c.alignment=R()
    if b: c2=ws.cell(r,2,b); c2.font=F(size=10); c2.alignment=R(w=True)
    r+=1
widths(ws,{'A':28,'B':110})

# ---------- Glossary
ws=wb.create_sheet(S_GL)
title(ws,'شرح المصطلحات بالعربي البسيط','لو أي كلمة في الملف مش واضحة، هتلاقي معناها هنا.')
gl=[('WSN','شبكة حساسات لاسلكية: أجهزة صغيرة ببطارية بتقيس حاجة (حرارة، رطوبة…) وتبعتها لجهاز رئيسي.'),
 ('BS / Sink','الجهاز الرئيسي اللي بيستلم الداتا من الشبكة كلها.'),
 ('CH (Cluster Head)','"الألفة": نود بيستلم داتا جيرانه، يجمّعها في رسالة واحدة، ويبعتها للـ BS. بيتعب أكتر من الباقيين.'),
 ('Cluster','مجموعة نودز تابعة لنفس الـ CH.'),
 ('Round','دورة واحدة: كل نود يبعت قراءته مرة، والـ CHs توصّلها للـ BS.'),
 ('Setup interval','كل كام round بنعيد اختيار الـ CHs من الأول (عندنا كل 5 rounds).'),
 ('FND (First Node Dead)','رقم الـ round اللي مات فيه أول نود. كل ما يكون أعلى، يبقى الحمل متوزع أحسن والأرض كلها متغطية أطول. (أعلى = أحسن)'),
 ('HND (Half Nodes Dead)','رقم الـ round اللي مات فيه نص النودز. بيقيس العمر المفيد للشبكة. (أعلى = أحسن)'),
 ('LND (Last Node Dead)','رقم الـ round اللي مات فيه آخر نود. ممكن يكون وهمي لو النودز عايشة بس مش بتبعت (زي v5b). (أعلى = أحسن، بس بحذر)'),
 ('PDR','نسبة الرسايل اللي وصلت للـ BS من كل الرسايل اللي اتعملت على طول عمر الشبكة. (أعلى = أحسن)'),
 ('Delivered packets','عدد الرسايل اللي وصلت فعلًا. بيكشف لو فيه نودز "عايشة وساكتة".'),
 ('Published','الأرقام زي ما هي مكتوبة في البحث الأصلي (كل بحث بإعدادات مختلفة).'),
 ('Paper-exact','الكود بتاعنا، بس بإعدادات البحث الأصلي نفسه — عشان نتأكد إن الكود صح.'),
 ('Unified / البيئة الموحّدة','الكود بتاعنا لكل البروتوكولات بنفس الإعدادات بالظبط — عشان المقارنة تبقى عادلة.'),
 ('Improved / بعد التحسين','البروتوكول بعد ما صلحنا فيه غلطة أو ضفنا تحسين.'),
 ('Δ (دلتا)','الفرق أو نسبة التغيير. +10% يعني زاد 10%، و‎−10% يعني قل 10%.'),
 ('pp (نقطة مئوية)','الفرق بين نسبتين: من 98.14% لـ 99.07% = +0.93 pp.'),
 ('FND ÷ LND','كل ما يقرّب من 1، يبقى النودز بتموت ورا بعض بسرعة، يعني الحمل كان متوزع بالعدل.'),
 ('Rank','الترتيب: 1 = الأحسن.'),
 ('Ablation','إننا نشيل آلية واحدة من البروتوكول ونشوف النتيجة وقعت قد إيه — عشان نعرف أهمية كل آلية.'),
 ('d² / d⁴','طاقة الإرسال بتزيد مع مربع المسافة في المسافات القصيرة، ومع أُس 4 للمسافة بعد 87.7 متر — عشان كده الإرسال البعيد غالي جدًا.')]
header(ws,4,['المصطلح','معناه ببساطة'])
for i,(a,b) in enumerate(gl): row(ws,5+i,[a,b],{},band=i%2==0)
widths(ws,{'A':26,'B':110})

# ---------- Unified
ws=wb.create_sheet(S_UNI)
title(ws,'المقارنة الموحّدة (الـ BS في نص الأرض)','كل الصفوف: 100 نود، 100×100 م، طاقة 0.5 جول، رسالة 2000 بت، نموذج راديو بمعادلتين، seed = 12345. الترتيب 1 = الأحسن.')
cols=['العيلة','البروتوكول / النسخة','المرحلة','الفكرة','FND','HND','LND','PDR','الـ LND مقارنة بـ LEACH','FND ÷ LND (توازن الحمل)','ترتيب LND','ترتيب PDR']
header(ws,4,cols,calc=(9,10,11,12))
rows=[(f,p,s,n,a,b,c,d) for f,p,s,a,b,c,d,n in D['unified']]
for v in D['proposed']:
    if v[1]=='في النص': rows.append(('مقترح',v[0],'مقترح',v[3],v[4],v[5],v[6],v[7]))
first=5; last=first+len(rows)-1
for i,rw in enumerate(rows):
    r=first+i
    row(ws,r,list(rw)+[f'=G{r}/$G${first}-1',f'=E{r}/G{r}',f'=RANK(G{r},$G${first}:$G${last})',f'=RANK(H{r},$H${first}:$H${last})'],{5:INT,6:INT,7:INT,8:PCT,9:DPCT,10:'0.00'},hl=rw[1].startswith('v8'),band=i%2==0)
note(ws,last+2,'إزاي تقري الجدول: "الـ LND مقارنة بـ LEACH" بتقارن بـ LEACH في البيئة الموحّدة. "FND ÷ LND" لو قريبة من 1 يبقى النودز ماتت ورا بعض بسرعة، يعني الحمل كان متوزع كويس.')
note(ws,last+3,'تنبيه: v5b عنده LND عالي (3203) بس ده وهمي — آخر 1038 round فيه وصّلوا 59 رسالة بس، لأن النودز كانت عايشة ومش بتبعت.')
ws.freeze_panes='C5'
widths(ws,{'A':16,'B':24,'C':17,'D':54,'E':9,'F':9,'G':9,'H':10,'I':14,'J':14,'K':10,'L':10})
UNI=(first,last)

# ---------- Reproduction
ws=wb.create_sheet(S_REP)
title(ws,'أرقام الأبحاث مقابل إعادة البناء بتاعتنا','المنشور = المكتوب في البحث (كل بحث بإعدادات مختلفة). Paper-exact = الكود بتاعنا بإعدادات البحث. الموحّد = الكود بتاعنا في البيئة المشتركة.')
cols=['البروتوكول','FND المنشور','LND المنشور','FND (paper-exact)','HND (paper-exact)','LND (paper-exact)','PDR (paper-exact)','FND الموحّد','HND الموحّد','LND الموحّد','PDR الموحّد','Paper-exact مقابل المنشور (LND)','الموحّد مقابل Paper-exact (LND)','إعدادات الـ paper-exact','مصدر الرقم المنشور']
header(ws,4,cols,calc=(12,13))
pub={p[0]:p for p in D['published']}; uni={}
for f,p,s,a,b,c,d,n in D['unified']:
    if s=='في البيئة الموحّدة' and p not in uni: uni[p]=(a,b,c,d)
for i,pe in enumerate(D['paperExact']):
    r=5+i; p=pe[0]; pb=pub[p]; u=uni[p]
    vals=[p,pb[1] if pb[1] else 'مش موجود',pb[3] if pb[3] else ('رسم بس' if p in('HEED','SH-LEACH') else 'مش موجود'),pe[1],pe[2],pe[3],pe[4],u[0],u[1],u[2],u[3],
          f'=IF(ISNUMBER(C{r}),F{r}/C{r}-1,"—")',f'=J{r}/F{r}-1',pe[5],pb[5]]
    row(ws,r,vals,{2:INT,3:INT,4:INT,5:INT,6:INT,7:PCT,8:INT,9:INT,10:INT,11:PCT,12:DPCT,13:DPCT},band=i%2==0)
n=5+len(D['paperExact'])+1
note(ws,n,'ليه الأرقام بتتغير؟ البيئة الموحّدة مختلفة في: مساحة الأرض، ومكان الـ BS، ونموذج الراديو، وحجم الرسالة، والطاقة الابتدائية. الخوارزمية نفسها ماتغيرتش.')
note(ws,n+1,'عشان كده عمود "الموحّد" هو الوحيد اللي ينفع نقارن بيه البروتوكولات ببعض بشكل عادل.')
ws.freeze_panes='B5'
widths(ws,{'A':12,'B':11,'C':11,'D':12,'E':12,'F':12,'G':12,'H':10,'I':10,'J':10,'K':10,'L':16,'M':16,'N':40,'O':44})

# ---------- Evolution
ws=wb.create_sheet(S_EVO)
title(ws,'البروتوكول المقترح — التطور نسخة بنسخة','كل نسخة بتغيّر حاجة واحدة بس. أعمدة Δ محسوبة: مقارنة بالصف اللي قبلها، ومقارنة بـ v1.')
cols=['النسخة','الـ BS','مدة إعادة التكوين','إيه اللي اتغيّر','FND','HND','LND','PDR','Δ FND عن اللي قبلها','Δ LND عن اللي قبلها','Δ PDR عن اللي قبلها','Δ LND عن v1']
header(ws,4,cols,calc=(9,10,11,12))
for i,v in enumerate(D['proposed']):
    r=5+i
    prev=[f'=E{r}/E{r-1}-1',f'=G{r}/G{r-1}-1',f'=(H{r}-H{r-1})*100'] if i else ['—','—','—']
    row(ws,r,[v[0],v[1],v[2],v[3],v[4],v[5],v[6],v[7]]+prev+[f'=G{r}/$G$5-1'],{5:INT,6:INT,7:INT,8:PCT,9:DPCT,10:DPCT,11:DPP,12:DPCT},hl=v[0].startswith('v8'),band=i%2==0)
EV=(5,4+len(D['proposed']))
note(ws,EV[1]+2,'الصفوف مترتبة بترتيب ما اتعملت. نسخ "في النص" و"بعيد" لنفس البروتوكول (v6 وv7 وv8) تتقارن مع بعض في شيت "نص مقابل بعيد"، مش صف بصف هنا.')
hb=EV[1]+5
ws.cell(hb-1,1,'بيانات الرسم (نسخ الـ BS في النص، بمعادلات)').font=F(bold=True,size=10,color=RED)
centre=[i for i,v in enumerate(D['proposed']) if v[1]=='في النص']
for k,t in enumerate(['النسخة','FND','HND','LND','PDR']): ws.cell(hb,k+1,t).font=F(bold=True,size=9)
for j,i in enumerate(centre):
    r=hb+1+j; src=5+i
    for k,col in enumerate('AEFGH'): ws.cell(r,k+1,f'={col}{src}').font=F(size=9)
    ws.cell(r,5).number_format=PCT
end=hb+len(centre)
ch=LineChart(); ch.title='عمر الشبكة عبر النسخ (BS في النص)'; ch.y_axis.title='Rounds'; ch.height=9; ch.width=24
ch.add_data(Reference(ws,min_col=2,max_col=4,min_row=hb,max_row=end),titles_from_data=True)
ch.set_categories(Reference(ws,min_col=1,min_row=hb+1,max_row=end))
for s,c in zip(ch.series,['C9A09E','8A2522','323232']): s.graphicalProperties.line.solidFill=c; s.graphicalProperties.line.width=28000
ws.add_chart(ch,f'G{hb-1}')
ws.freeze_panes='B5'
widths(ws,{'A':24,'B':9,'C':11,'D':56,'E':9,'F':9,'G':9,'H':10,'I':13,'J':13,'K':13,'L':12})

# ---------- Center vs Far
ws=wb.create_sheet(S_CF)
title(ws,'نفس البروتوكول: الـ BS في النص مقابل الـ BS بعيد','نفس الكود ونفس الـ seed، الفرق الوحيد مكان الـ BS: في النص (50,50) أو بعيد (50,−100). Δ = البعيد مقارنة بالنص (معادلات).')
cols=['البروتوكول','FND نص','HND نص','LND نص','PDR نص','FND بعيد','HND بعيد','LND بعيد','PDR بعيد','Δ FND','Δ LND','Δ PDR']
header(ws,4,cols,calc=(10,11,12))
pairs=[('v5b (Energy-Aware)',(1730,2091,3203,0.973087),(1055,1616,1861,0.982785)),('v6 (سلسلة ثابتة)',(1551,1981,2106,0.97494),(1266,1910,2031,0.97075)),
       ('v7 (سلسلة + Backup + إصلاح)',(1551,1976,2066,0.98432),(1266,1891,1976,0.98135)),('v8',(2621,2711,2761,0.992224),(1471,1646,1716,0.986599)),
       ('v8-Chain (توجيه ذكي)',(2621,2711,2756,0.992668),(1716,1886,1956,0.990719))]
for i,(p,c,f) in enumerate(pairs):
    r=5+i
    row(ws,r,[p,*c,*f,f'=F{r}/B{r}-1',f'=H{r}/D{r}-1',f'=(I{r}-E{r})*100'],{2:INT,3:INT,4:INT,5:PCT,6:INT,7:INT,8:INT,9:PCT,10:DPCT,11:DPCT,12:DPP},hl=p.startswith('v8-Chain'),band=i%2==0)
note(ws,11,'إزاي تقري الجدول: لما الـ BS يبعد، كل البروتوكولات بتخسر عمر لأن الإرسال للـ BS بقى غالي. v8-Chain بيخسر أقل من v8 لأنه بيعمل سلسلة لوحده لما الـ BS يبعد.')
ch=BarChart(); ch.type='col'; ch.grouping='clustered'; ch.title='العمر لما الـ BS بعيد'; ch.y_axis.title='Rounds'; ch.height=8.5; ch.width=20
ch.add_data(Reference(ws,min_col=6,max_col=8,min_row=4,max_row=9),titles_from_data=True)
ch.set_categories(Reference(ws,min_col=1,min_row=5,max_row=9))
for s,c in zip(ch.series,['C9A09E','8A2522','323232']): s.graphicalProperties.solidFill=c; s.graphicalProperties.line.solidFill=c
ws.add_chart(ch,'A13')
widths(ws,{'A':30,'B':10,'C':10,'D':10,'E':10,'F':10,'G':10,'H':10,'I':10,'J':10,'K':10,'L':11})

# ---------- v8 analysis
ws=wb.create_sheet(S_V8)
title(ws,'تحليل v8 / v8-Chain','أ. نشيل آلية واحدة كل مرة (Ablation). ب. طرق التوجيه بين الـ CHs. ج. متوسط 8 أشكال شبكات مختلفة.')
ws['A4']='أ. Ablation — نشيل آلية واحدة (الـ BS في النص، seed 12345)'; ws['A4'].font=F(bold=True,size=11,color=RED); ws['A4'].alignment=R()
header(ws,5,['النسخة','الآلية اللي اتشالت','FND','HND','LND','PDR','الرسايل اللي وصلت','Δ FND عن الكامل','Δ الرسايل عن الكامل'],calc=(8,9))
for i,a in enumerate(D['ablation']):
    r=6+i; row(ws,r,a+[f'=C{r}/$C$6-1',f'=G{r}/$G$6-1'],{3:INT,4:INT,5:INT,6:PCT,7:INT,8:DPCT,9:DPCT},hl=i==0,band=i%2==0)
r0=6+len(D['ablation'])+2
ws.cell(r0,1,'ب. طريقة التوجيه بين الـ CHs (seed 12345 ومتوسط 8 seeds)').font=F(bold=True,size=11,color=RED); ws.cell(r0,1).alignment=R()
header(ws,r0+1,['طريقة التوجيه','الـ BS','FND','HND','LND','PDR','متوسط FND','متوسط HND','متوسط LND','متوسط PDR'])
for i,x in enumerate(D['routing']): row(ws,r0+2+i,x,{3:INT,4:INT,5:INT,6:PCT,7:INT,8:INT,9:INT,10:PCT},hl=x[0].startswith('توجيه ذكي'),band=i%2==0)
r1=r0+2+len(D['routing'])+2
ws.cell(r1,1,'ج. الثبات — متوسط 8 أشكال شبكات (seeds 12345, 1, 7, 42, 99, 2024, 31337, 555)').font=F(bold=True,size=11,color=RED); ws.cell(r1,1).alignment=R()
header(ws,r1+1,['البروتوكول','الـ BS','متوسط FND','متوسط HND','متوسط LND','متوسط PDR'])
for i,x in enumerate(D['robust']): row(ws,r1+2+i,x,{3:INT,4:INT,5:INT,6:PCT},hl=x[0]=='v8',band=i%2==0)
k=r1+2+len(D['robust'])+1
note(ws,k,'الخلاصة: الإرسال المباشر للـ BS (I4) وترشيح الأقوياء بس (I5) هما سبب معظم الزيادة في العمر. التسليم الاستباقي (I2) بيزوّد الـ PDR.')
note(ws,k+1,'I1 وI3 شبكة أمان: مالهمش أثر لما كل الآليات شغّالة مع بعض في الـ seed ده. التوجيه الذكي مابيشتغلش خالص لما الـ BS في النص، وبيعمل سلسلة لوحده لما الـ BS بعيد.')
widths(ws,{'A':32,'B':26,'C':10,'D':10,'E':10,'F':10,'G':16,'H':14,'I':16,'J':10})

# ---------- Dashboard
ws=wb.create_sheet(S_DASH,1)
title(ws,'لوحة الملخص — أحسن نسخة من كل عيلة (الـ BS في النص)',f'الأرقام بتتسحب من شيت "{S_UNI}" بـ INDEX/MATCH، فلو غيّرتي رقم هناك بيتغير هنا لوحده.')
pick=['LEACH','HEED','PEGASIS','SH-LEACH','H-LEACH','v3','v7 (Center)','v7.1 Multihop (Int 5)','v8-Chain (Center)']
stage={'SH-LEACH':'بعد التحسين','H-LEACH':'بعد التحسين'}
header(ws,4,['البروتوكول','المرحلة','FND','HND','LND','PDR'],calc=(3,4,5,6))
U=q(S_UNI)
for i,p in enumerate(pick):
    r=5+i; st=stage.get(p,'في البيئة الموحّدة' if not p.startswith('v') else 'مقترح')
    ws.cell(r,1,p); ws.cell(r,2,st)
    key=f'MATCH(1,INDEX(({U}!$B${UNI[0]}:$B${UNI[1]}=A{r})*({U}!$C${UNI[0]}:$C${UNI[1]}=B{r}),0),0)'
    for j,col in zip(range(3,7),'EFGH'): ws.cell(r,j,f'=INDEX({U}!${col}${UNI[0]}:${col}${UNI[1]},{key})')
    for j in range(1,7):
        c=ws.cell(r,j); c.font=F(size=10,bold=p.startswith('v8')); c.border=B; c.fill=HL if p.startswith('v8') else (BAND if i%2==0 else NOF)
        c.number_format=PCT if j==6 else INT; c.alignment=R(h='right' if j<3 else 'center')
lastd=5+len(pick)-1
ws.cell(lastd+2,1,'أهم الأرقام').font=F(bold=True,size=11,color=RED)
C2=q(S_CF)
kp=[('FND بتاع v8-Chain مقارنة بأحسن هجين من الأبحاث (SH-LEACH المحسّن)',f'=C{lastd}/C8-1',DPCT),
    ('FND بتاع v8-Chain مقارنة بـ LEACH',f'=C{lastd}/C5-1',DPCT),
    ('PDR بتاع v8-Chain مقارنة بـ v7 (نص)',f'=(F{lastd}-F11)*100',DPP),
    ('FND بتاع v8-Chain لما الـ BS بعيد مقارنة بـ v7 بعيد',f"={C2}!F9/{C2}!F7-1",DPCT)]
for k,(a,f,fm) in enumerate(kp):
    r=lastd+3+k; ws.cell(r,1,a).font=F(size=10); ws.cell(r,1).alignment=R(); c=ws.cell(r,3,f); c.number_format=fm; c.font=F(bold=True,size=12,color=RED)
note(ws,lastd+8,'ملاحظة مهمة: LEACH وPEGASIS وHEED وH-LEACH المحسّن عندهم PDR أعلى شوية من v8-Chain (99.36–99.68%)، بس أول نود فيهم بتموت قبلنا بـ 1,100 لـ 2,000 round. وPEGASIS عنده LND أعلى (3504)، بس أول نود فيه بتموت قبلنا بحوالي 1300 round.')
ch=BarChart(); ch.type='col'; ch.grouping='clustered'; ch.title='عمر الشبكة (rounds)'; ch.height=9; ch.width=22
ch.add_data(Reference(ws,min_col=3,max_col=5,min_row=4,max_row=lastd),titles_from_data=True)
ch.set_categories(Reference(ws,min_col=1,min_row=5,max_row=lastd))
for s,c in zip(ch.series,['C9A09E','8A2522','323232']): s.graphicalProperties.solidFill=c; s.graphicalProperties.line.solidFill=c
ws.add_chart(ch,'H4')
ch2=BarChart(); ch2.type='col'; ch2.title='نسبة التوصيل PDR'; ch2.height=7.5; ch2.width=22; ch2.legend=None
ch2.add_data(Reference(ws,min_col=6,max_col=6,min_row=4,max_row=lastd),titles_from_data=True)
ch2.set_categories(Reference(ws,min_col=1,min_row=5,max_row=lastd))
ch2.y_axis.scaling.min=0.96; ch2.y_axis.scaling.max=1.0; ch2.y_axis.number_format='0.0%'
ch2.series[0].graphicalProperties.solidFill='8A2522'
ch2.dataLabels=DataLabelList(); ch2.dataLabels.showVal=True; ch2.dataLabels.showSerName=False; ch2.dataLabels.showCatName=False; ch2.dataLabels.showLegendKey=False; ch2.dataLabels.numFmt='0.00%'
ws.add_chart(ch2,'H23')
widths(ws,{'A':46,'B':17,'C':10,'D':10,'E':10,'F':10})
for w in wb.worksheets:
    w.page_setup.orientation='landscape'; w.sheet_properties.pageSetUpPr.fitToPage=True; w.page_setup.fitToWidth=1; w.page_setup.fitToHeight=0
wb.save('WSN_Protocol_Comparison_AR.xlsx'); print('saved')
